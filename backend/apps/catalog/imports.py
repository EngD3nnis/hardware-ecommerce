"""Spreadsheet (xlsx/csv) product import: validate → approve → commit → (rollback).

Workflow (driven from the admin, run in Celery):
1. Staff upload a file → ImportBatch(UPLOADED).
2. validate_batch(): parse, normalise and check every row; nothing is written
   to products. Each row gets an outcome, and the batch gets a report.
3. A person reads the report and approves → commit_batch() applies every
   accepted row in ONE transaction: all rows or none.
4. rollback_batch() undoes a committed batch as far as is safe: new products
   are deleted (or archived if already referenced), updated fields are
   restored unless someone changed them since, and prices it set are removed.

Recognised columns (headers are case-insensitive; see COLUMN_ALIASES):
sku, name, category, subcategory, brand, unit, barcode, description,
material, finish, colour, size, price, plus attribute columns named
"attr:<code>" or matching an attribute's code or label.
"""

import csv
import io
import re
from collections import Counter
from decimal import Decimal

from django.db import transaction
from django.db.models import ProtectedError
from django.utils import timezone

from apps.audit import services as audit
from apps.core.actors import Actor
from apps.core.exceptions import Conflict, DomainError, ValidationError
from apps.pricing import services as pricing
from apps.pricing.models import ProductPrice

from . import services
from .media import sha256_hex
from .models import (
    SKU_PATTERN,
    AttributeDefinition,
    Brand,
    Category,
    ImportBatch,
    ImportRow,
    ImportStatus,
    Product,
    ProductStatus,
    RowOutcome,
    UnitOfMeasure,
    normalise_sku,
)

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_ROWS = 20_000
REPORT_EXAMPLES = 50

COLUMN_ALIASES = {
    "sku": ["sku", "item code", "itemcode", "code", "product code", "item no", "item number", "part no"],
    "name": ["name", "product name", "item name", "product", "item", "item description"],
    "category": ["category", "cat", "department", "group"],
    "subcategory": ["subcategory", "sub category", "sub-category", "sub", "sub group", "subgroup"],
    "brand": ["brand", "make", "manufacturer"],
    "unit": ["unit", "uom", "unit of measure", "units"],
    "barcode": ["barcode", "ean", "gtin", "upc"],
    "description": ["description", "details", "long description"],
    "material": ["material"],
    "finish": ["finish"],
    "colour": ["colour", "color"],
    "size": ["size", "size label"],
    "price": ["price", "selling price", "retail price", "unit price", "price (kes)", "price kes", "sp"],
}

UNIT_ALIASES = {
    UnitOfMeasure.PIECE: ["pc", "pcs", "piece", "pieces", "each", "ea", "no", "nos", "unit", "units"],
    UnitOfMeasure.PAIR: ["pair", "pairs", "pr"],
    UnitOfMeasure.SET: ["set", "sets", "kit"],
    UnitOfMeasure.METRE: ["m", "metre", "metres", "meter", "meters", "mtr", "mtrs"],
    UnitOfMeasure.LENGTH: ["length", "lengths", "len", "lgth"],
    UnitOfMeasure.ROLL: ["roll", "rolls", "rl"],
    UnitOfMeasure.SHEET: ["sheet", "sheets", "sht"],
    UnitOfMeasure.KG: ["kg", "kgs", "kilogram", "kilograms"],
    UnitOfMeasure.BAG: ["bag", "bags"],
    UnitOfMeasure.LITRE: ["l", "ltr", "ltrs", "litre", "litres", "liter", "liters"],
    UnitOfMeasure.TIN: ["tin", "tins", "can"],
    UnitOfMeasure.BOX: ["box", "boxes", "ctn", "carton", "cartons"],
    UnitOfMeasure.PACKET: ["pkt", "pkts", "packet", "packets", "pack", "packs"],
}
_UNIT_LOOKUP = {alias: unit for unit, aliases in UNIT_ALIASES.items() for alias in aliases}
_UNIT_LOOKUP.update({u.value.lower(): u for u in UnitOfMeasure})

SIMPLE_TEXT_FIELDS = ("description", "material", "finish", "colour")
SUSPICIOUS_PRICE_LOW = Decimal("1")
SUSPICIOUS_PRICE_HIGH = Decimal("1000000")


# --- Parsing -------------------------------------------------------------------------


def _cell_to_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # Excel stores 1001 and barcodes as floats
    return " ".join(str(value).split())


def read_table(data: bytes, filename: str) -> tuple[list[str], list[tuple[int, list[str]]]]:
    """Return (headers, [(row_number, cells)]) from an .xlsx or .csv file."""
    name = filename.lower()
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        try:
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001 - openpyxl raises many types for bad files
            raise ValidationError("The file is not a readable .xlsx workbook.") from exc
        sheet = workbook.worksheets[0]
        rows = [(i, [_cell_to_text(c) for c in row]) for i, row in enumerate(sheet.iter_rows(values_only=True), 1)]
        workbook.close()
    elif name.endswith(".csv"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("latin-1")
        rows = [(i, [_cell_to_text(c) for c in row]) for i, row in enumerate(csv.reader(io.StringIO(text)), 1)]
    else:
        raise ValidationError("Upload an .xlsx or .csv file.")

    rows = [(n, cells) for n, cells in rows if any(cells)]
    if not rows:
        raise ValidationError("The file is empty.")
    header_number, headers = rows[0]
    body = rows[1:]
    if len(body) > MAX_ROWS:
        raise ValidationError(f"Too many rows ({len(body)}); split the file into parts of at most {MAX_ROWS}.")
    return headers, body


def map_columns(headers: list[str]) -> tuple[dict[str, int], dict[str, int], list[str]]:
    """Return (field → column index, attribute code → column index, unrecognised headers)."""
    alias_to_field = {alias: f for f, aliases in COLUMN_ALIASES.items() for alias in aliases}
    attributes = {a.code.lower(): a.code for a in AttributeDefinition.objects.all()}
    attributes.update({a.label.lower(): a.code for a in AttributeDefinition.objects.all()})
    fields, attr_columns, unknown = {}, {}, []
    for index, header in enumerate(headers):
        key = header.strip().lower()
        if not key:
            continue
        if key.startswith("attr:"):
            attr_columns[key[5:].strip()] = index
        elif key in alias_to_field and alias_to_field[key] not in fields:
            fields[alias_to_field[key]] = index
        elif key in attributes:
            attr_columns[attributes[key]] = index
        else:
            unknown.append(header)
    return fields, attr_columns, unknown


# --- Validation ----------------------------------------------------------------------


def _find_category(category: str, subcategory: str) -> Category | None:
    top = Category.objects.filter(parent__isnull=True, name__iexact=category).first()
    if top is None or not subcategory or subcategory.lower() == category.lower():
        return top
    return Category.objects.filter(parent=top, name__iexact=subcategory).first()


def _normalise_row(cells, fields, attr_columns, definitions) -> tuple[dict, list[str], list[str]]:
    def cell(field):
        index = fields.get(field)
        return cells[index].strip() if index is not None and index < len(cells) else ""

    errors, warnings = [], []
    row = {
        "sku": normalise_sku(cell("sku")),
        "name": services.clean_name(cell("name")),
        "category": services.clean_name(cell("category")),
        "subcategory": services.clean_name(cell("subcategory")),
        "brand": services.clean_name(cell("brand")),
        "barcode": cell("barcode").replace(" ", ""),
        "size_label": cell("size"),
        "attributes": {},
    }
    for field in SIMPLE_TEXT_FIELDS:
        row[field] = cell(field)

    if not row["sku"]:
        errors.append("Missing SKU.")
    elif not SKU_PATTERN.match(row["sku"]):
        errors.append(f"Invalid SKU {row['sku']!r} (letters, digits and . _ / - only).")
    if not row["name"]:
        errors.append("Missing product name.")
    elif len(row["name"]) < 3 or row["name"].isdigit():
        errors.append(f"Suspicious product name {row['name']!r}.")
    elif len(row["name"]) > 200:
        warnings.append("Suspicious value: name is longer than 200 characters.")
    if not row["category"]:
        errors.append("Missing category.")
    if row["barcode"] and not re.fullmatch(r"\d{8,14}", row["barcode"]):
        errors.append(f"Invalid barcode {row['barcode']!r} (8-14 digits).")

    raw_unit = cell("unit")
    unit = _UNIT_LOOKUP.get(raw_unit.lower().rstrip(".")) if raw_unit else UnitOfMeasure.PIECE
    if unit is None:
        warnings.append(f"Unknown unit {raw_unit!r}; Piece will be used.")
        unit = UnitOfMeasure.PIECE
    row["unit_of_measure"] = UnitOfMeasure(unit).value
    row["unit_given"] = bool(raw_unit) and unit is not None

    raw_price = cell("price")
    row["price"] = None
    if raw_price:
        try:
            price = pricing.parse_amount(raw_price)
            row["price"] = str(price)
            if price < SUSPICIOUS_PRICE_LOW or price > SUSPICIOUS_PRICE_HIGH:
                warnings.append(f"Suspicious value: price {price}.")
        except ValidationError as exc:
            errors.append(f"{exc.message} ({raw_price!r})")

    for code, index in attr_columns.items():
        raw = cells[index].strip() if index < len(cells) else ""
        if not raw:
            continue
        definition = definitions.get(code)
        if definition is None:
            warnings.append(f"Unknown attribute column {code!r} ignored.")
            continue
        try:
            services.coerce_attribute_value(definition, raw)
            row["attributes"][code] = raw
        except ValidationError as exc:
            errors.append(exc.message)
    return row, errors, warnings


def _product_changes(product: Product, row: dict, category: Category | None) -> dict:
    """Plain-value differences between an existing product and a row (blank cells never clear data)."""
    changes = {}
    if row["name"] and row["name"] != product.name:
        changes["name"] = row["name"]
    if category is not None and category.pk != product.category_id:
        changes["category"] = category.slug
    if row["brand"] and (product.brand is None or product.brand.name.lower() != row["brand"].lower()):
        changes["brand"] = row["brand"]
    if row["barcode"] and row["barcode"] != (product.barcode or ""):
        changes["barcode"] = row["barcode"]
    if row["size_label"] and row["size_label"] != product.size_label:
        changes["size_label"] = row["size_label"]
    for field in SIMPLE_TEXT_FIELDS:
        if row[field] and row[field] != getattr(product, field):
            changes[field] = row[field]
    if row["unit_given"] and row["unit_of_measure"] != product.unit_of_measure:
        changes["unit_of_measure"] = row["unit_of_measure"]
    return changes


def validate_batch(batch: ImportBatch) -> ImportBatch:
    """Dry run: classify every row and build the report. Writes only ImportRows."""
    with transaction.atomic():
        batch = ImportBatch.objects.select_for_update().get(pk=batch.pk)
        if batch.status not in (ImportStatus.UPLOADED, ImportStatus.VALIDATED, ImportStatus.FAILED):
            raise Conflict(f"Cannot validate a batch that is {batch.get_status_display().lower()}.")
        batch.status = ImportStatus.VALIDATING
        batch.save(update_fields=["status", "updated_at"])

    try:
        batch.file.open("rb")
        data = batch.file.read()
        batch.file.close()
        if len(data) > MAX_FILE_BYTES:
            raise ValidationError("File is larger than 10 MB.")
        headers, body = read_table(data, batch.original_filename or batch.file.name)
        fields, attr_columns, unknown = map_columns(headers)
        missing = [f for f in ("sku", "name", "category") if f not in fields]
        if missing:
            raise ValidationError(f"Required column(s) not found: {', '.join(missing)}. Headers seen: {headers}")
    except ValidationError as exc:
        batch.status, batch.error = ImportStatus.FAILED, exc.message
        batch.save(update_fields=["status", "error", "updated_at"])
        return batch

    definitions = {a.code: a for a in AttributeDefinition.objects.filter(code__in=attr_columns)}
    existing = {
        p.sku: p
        for p in Product.objects.filter(
            sku__in=[normalise_sku(c[fields["sku"]]) for _, c in body if fields["sku"] < len(c)]
        ).select_related("brand")
    }
    names_in_db = dict(Product.objects.exclude(status=ProductStatus.ARCHIVED).values_list("name", "sku"))
    names_in_db = {name.lower(): sku for name, sku in names_in_db.items()}
    barcodes_in_db = dict(Product.objects.exclude(barcode__isnull=True).values_list("barcode", "sku"))

    rows, first_seen = [], {}
    new_categories = set()
    for number, cells in body:
        raw = {headers[i] or f"column {i + 1}": v for i, v in enumerate(cells) if i < len(headers) and v}
        normalized, errors, warnings = _normalise_row(cells, fields, attr_columns, definitions)
        sku = normalized["sku"]
        outcome = RowOutcome.REJECTED if errors else None

        if not errors and sku in first_seen:
            outcome = RowOutcome.DUPLICATE
            errors.append(f"SKU {sku} already appears on row {first_seen[sku]}.")
        if sku and sku not in first_seen:
            first_seen[sku] = number

        category = None
        if not errors:
            category = _find_category(normalized["category"], normalized["subcategory"])
            if category is None:
                label = " / ".join(x for x in (normalized["category"], normalized["subcategory"]) if x)
                new_categories.add(label)
                warnings.append(f"New category will be created: {label}.")
            owner = barcodes_in_db.get(normalized["barcode"]) if normalized["barcode"] else None
            if owner and owner != sku:
                errors.append(f"Barcode {normalized['barcode']} already belongs to {owner}.")
            for code in normalized["attributes"]:
                try:
                    services.check_attribute_scope(definitions[code], category)
                except ValidationError as exc:
                    errors.append(exc.message)
            if errors:
                outcome = RowOutcome.REJECTED

        if not errors:
            product = existing.get(sku)
            if product is None:
                outcome = RowOutcome.CREATE
                dup_sku = names_in_db.get(normalized["name"].lower())
                if dup_sku:
                    warnings.append(f"Possible duplicate: {dup_sku} has the same name.")
            else:
                changes = _product_changes(product, normalized, category)
                normalized["changes"] = changes
                price_changed = normalized["price"] is not None and _price_differs(product, normalized["price"])
                if not changes and not price_changed and not normalized["attributes"]:
                    outcome = RowOutcome.UNCHANGED
                elif batch.update_existing:
                    outcome = RowOutcome.UPDATE
                else:
                    outcome = RowOutcome.PROPOSE
                    if price_changed:
                        warnings.append(
                            "Price differs from the current price; not changed (proposals cannot set prices)."
                        )

        rows.append(
            ImportRow(
                batch=batch,
                row_number=number,
                raw=raw,
                normalized=normalized,
                outcome=outcome,
                errors=errors,
                warnings=warnings,
            )
        )

    counts = Counter(r.outcome for r in rows)
    report = {
        "rows_processed": len(rows),
        "rows_accepted": sum(
            counts[o] for o in (RowOutcome.CREATE, RowOutcome.UPDATE, RowOutcome.PROPOSE, RowOutcome.UNCHANGED)
        ),
        "rows_rejected": counts[RowOutcome.REJECTED],
        "duplicates": counts[RowOutcome.DUPLICATE],
        "outcomes": {o.label: counts[o] for o in RowOutcome if counts[o]},
        "missing_skus": sum(1 for r in rows if "Missing SKU." in r.errors),
        "missing_categories": sum(1 for r in rows if "Missing category." in r.errors),
        "new_categories": sorted(new_categories),
        "missing_images": counts[RowOutcome.CREATE],  # new products never have images yet
        "suspicious_values": sum(1 for r in rows for w in r.warnings + r.errors if "Suspicious" in w),
        "possible_duplicates": sum(1 for r in rows for w in r.warnings if w.startswith("Possible duplicate")),
        "unrecognised_columns": unknown,
        "attribute_columns": sorted(attr_columns),
        "examples": [
            {"row": r.row_number, "outcome": r.outcome, "errors": r.errors, "warnings": r.warnings}
            for r in rows
            if r.errors or r.warnings
        ][:REPORT_EXAMPLES],
    }
    with transaction.atomic():
        ImportRow.objects.filter(batch=batch).delete()
        ImportRow.objects.bulk_create(rows, batch_size=1000)
        batch.report, batch.error, batch.status = report, "", ImportStatus.VALIDATED
        batch.save(update_fields=["report", "error", "status", "updated_at"])
    return batch


def _price_differs(product: Product, amount: str) -> bool:
    current = pricing.current_price(product)
    return current is None or current.amount != Decimal(amount)


# --- Commit --------------------------------------------------------------------------------


def _resolve_category(row: dict, actor: Actor) -> Category:
    top, _ = services.get_or_create_category(row["category"], actor)
    sub = row["subcategory"]
    if not sub or sub.lower() == top.name.lower():
        return top
    return services.get_or_create_category(sub, actor, parent=top)[0]


def _apply_price(row: ImportRow, product: Product, actor: Actor, batch: ImportBatch) -> None:
    if row.normalized.get("price") is None:
        return
    price = pricing.set_price(product, row.normalized["price"], actor, source=f"import:{batch.pk}")
    if price is not None:
        row.created_price_id = price.pk
        if not product.price_on_request:
            return
        # A list price was supplied: the product no longer needs "price on request".
        services.update_product(product, {"price_on_request": False}, actor, reason=f"Import #{batch.pk}")


def commit_batch(batch: ImportBatch, actor: Actor) -> ImportBatch:
    """Apply all accepted rows atomically. Idempotent: a committed batch is left alone."""
    with transaction.atomic():
        batch = ImportBatch.objects.select_for_update().get(pk=batch.pk)
        if batch.status == ImportStatus.COMMITTED:
            return batch
        if batch.status != ImportStatus.VALIDATED:
            raise Conflict(f"Only a validated batch can be imported (this one is {batch.get_status_display()}).")
        batch.status = ImportStatus.COMMITTING
        batch.save(update_fields=["status", "updated_at"])

    try:
        with transaction.atomic():
            _commit_rows(batch, actor)
            batch.status, batch.approved_by, batch.committed_at = ImportStatus.COMMITTED, actor.label, timezone.now()
            batch.error = ""
            batch.save(update_fields=["status", "approved_by", "committed_at", "error", "updated_at"])
            audit.record(actor, "catalog.import.committed", batch, after=batch.report.get("outcomes"))
    except DomainError as exc:
        # Nothing was applied (single transaction). Data may have changed since validation.
        ImportBatch.objects.filter(pk=batch.pk).update(
            status=ImportStatus.FAILED, error=f"{exc.message} Re-validate the batch and try again."
        )
        batch.refresh_from_db()
    return batch


def _commit_rows(batch: ImportBatch, actor: Actor) -> None:
    status = ProductStatus.ACTIVE if batch.activate_new else ProductStatus.DRAFT
    for row in batch.rows.exclude(outcome__in=[RowOutcome.REJECTED, RowOutcome.DUPLICATE, RowOutcome.UNCHANGED]):
        data = row.normalized
        if row.outcome == RowOutcome.CREATE:
            brand = services.get_or_create_brand(data["brand"], actor)[0] if data["brand"] else None
            product = services.create_product(
                actor,
                sku=data["sku"],
                name=data["name"],
                category=_resolve_category(data, actor),
                brand=brand,
                status=status,
                unit_of_measure=data["unit_of_measure"],
                barcode=data["barcode"] or None,
                size_label=data["size_label"],
                **{f: data[f] for f in SIMPLE_TEXT_FIELDS},
            )
            row.product = product
        else:
            product = Product.objects.get(sku=data["sku"])
            row.product = product
            changes = data.get("changes", {})
            if row.outcome == RowOutcome.UPDATE and changes:
                resolved = dict(changes)
                if "category" in resolved:
                    resolved["category"] = _resolve_category(data, actor)
                if "brand" in resolved:
                    resolved["brand"] = services.get_or_create_brand(resolved["brand"], actor)[0]
                if "barcode" in resolved:
                    resolved["barcode"] = resolved["barcode"] or None
                row.before = audit.snapshot(product, list(resolved))
                services.update_product(product, resolved, actor, reason=f"Import #{batch.pk}, row {row.row_number}")
                row.after = audit.snapshot(product, list(resolved))
            elif row.outcome == RowOutcome.PROPOSE:
                for field, value in changes.items():
                    if field == "category" and Category.objects.filter(slug=value).first() is None:
                        continue  # a new category must be created by a person first
                    services.propose_change(
                        product, field, value, actor, source=f"import:{batch.pk}", reason=f"Row {row.row_number}"
                    )
        for code, raw in data.get("attributes", {}).items():
            services.set_attribute(product, code, raw, actor)
        if row.outcome != RowOutcome.PROPOSE:
            _apply_price(row, product, actor, batch)
        row.applied = True
        row.save(update_fields=["product", "applied", "before", "after", "created_price_id"])


# --- Rollback -------------------------------------------------------------------------------


@transaction.atomic
def rollback_batch(batch: ImportBatch, actor: Actor) -> dict:
    """Undo a committed batch where it is still safe. Returns what was and wasn't reverted."""
    batch = ImportBatch.objects.select_for_update().get(pk=batch.pk)
    if batch.status != ImportStatus.COMMITTED:
        raise Conflict("Only an imported batch can be rolled back.")
    result = Counter()
    skipped = []
    for row in batch.rows.filter(applied=True).select_related("product").order_by("-row_number"):
        product = row.product
        if row.created_price_id:
            _rollback_price(row.created_price_id, result, skipped, row)
        if product is None:
            continue
        if row.outcome == RowOutcome.CREATE:
            sku = product.sku
            try:
                with transaction.atomic():
                    product.delete()
                result["products_deleted"] += 1
                audit.record(
                    actor,
                    "catalog.product.deleted",
                    object_type="catalog.product",
                    object_id=sku,
                    reason=f"Rollback of import #{batch.pk}",
                )
            except ProtectedError:
                services.set_status(product, ProductStatus.ARCHIVED, actor, reason=f"Rollback of import #{batch.pk}")
                result["products_archived"] += 1
        elif row.outcome == RowOutcome.UPDATE and row.before:
            current = audit.snapshot(product, list(row.before))
            if current != row.after:
                skipped.append({"row": row.row_number, "sku": product.sku, "reason": "changed since import"})
                continue
            restore = {}
            for field, value in row.before.items():
                if field in ("category", "brand"):
                    model = Category if field == "category" else Brand
                    restore[field] = model.objects.filter(pk=value).first() if value else None
                else:
                    restore[field] = value
            services.update_product(product, restore, actor, reason=f"Rollback of import #{batch.pk}")
            result["products_restored"] += 1
    batch.status, batch.rolled_back_by, batch.rolled_back_at = ImportStatus.ROLLED_BACK, actor.label, timezone.now()
    batch.report = {**batch.report, "rollback": {**result, "skipped": skipped}}
    batch.save(update_fields=["status", "rolled_back_by", "rolled_back_at", "report", "updated_at"])
    audit.record(actor, "catalog.import.rolled_back", batch, after={**result, "skipped": len(skipped)})
    return {**result, "skipped": skipped}


def _rollback_price(price_id: int, result: Counter, skipped: list, row: ImportRow) -> None:
    price = ProductPrice.objects.filter(pk=price_id).first()
    if price is None:
        return
    if price.valid_to is not None:
        skipped.append({"row": row.row_number, "reason": "price already replaced since import"})
        return
    previous = ProductPrice.objects.filter(
        product=price.product, price_list=price.price_list, valid_to=price.valid_from
    ).first()
    price.delete()
    if previous:
        previous.valid_to = None
        previous.save(update_fields=["valid_to", "updated_at"])
    result["prices_removed"] += 1


def create_batch(uploaded_file, actor: Actor, *, update_existing=False, activate_new=False) -> ImportBatch:
    data = uploaded_file.read()
    if len(data) > MAX_FILE_BYTES:
        raise ValidationError("File is larger than 10 MB.")
    uploaded_file.seek(0)
    return ImportBatch.objects.create(
        file=uploaded_file,
        original_filename=getattr(uploaded_file, "name", "")[:255],
        file_sha256=sha256_hex(data),
        update_existing=update_existing,
        activate_new=activate_new,
        created_by=actor.label,
    )
